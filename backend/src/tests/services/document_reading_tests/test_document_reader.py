import os
import zipfile

import pytest

from api.services.document_reading import FINAL_STATUSES, DocumentReadLimits, is_document_path, read_document_text

from .document_fixtures import (
    make_docx_package,
    make_iwork,
    make_pdf,
    make_pptx,
    make_xlsx,
    pdf_bytes,
    word_document,
    write_zip,
)


def test_pdf_pages_are_labeled_and_text_is_extracted(tmp_path):
    path = make_pdf(tmp_path / "plan.pdf", ["Launch retry policy review", "Second page budget"])
    result = read_document_text(path)
    assert result.status == "ok" and result.kind == "pdf" and result.method == "pypdf"
    assert "[Page 1]\nLaunch retry policy review" in result.text
    assert "[Page 2]\nSecond page budget" in result.text
    assert result.units == 2 and result.truncated is False


def test_pdf_without_text_reports_no_text(tmp_path):
    result = read_document_text(make_pdf(tmp_path / "scan.pdf", [""]))
    assert result.status == "no_text" and result.text == "" and "OCR" in result.detail


def test_pdf_page_limit_truncates_with_marker(tmp_path):
    path = make_pdf(tmp_path / "long.pdf", ["alpha page", "beta page", "gamma page"])
    result = read_document_text(path, DocumentReadLimits(max_pdf_pages=1))
    assert result.truncated is True and result.status == "ok"
    assert "alpha page" in result.text and "beta page" not in result.text
    assert result.text.endswith("[Text truncated at 400000 characters]")


def test_encrypted_pdf_reports_encrypted(tmp_path):
    from pypdf import PdfReader, PdfWriter
    import io

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes(["secret text"]))))
    writer.encrypt("hunter2", algorithm="RC4-128")
    path = tmp_path / "locked.pdf"
    with open(path, "wb") as handle:
        writer.write(handle)
    result = read_document_text(path)
    assert result.status == "encrypted" and result.text == ""


def test_pdf_deadline_returns_partial_text_that_is_not_final(tmp_path):
    ticks = iter(range(0, 1000))
    path = make_pdf(tmp_path / "slow.pdf", ["one", "two", "three", "four", "five"])
    result = read_document_text(path, deadline_seconds=2.5, clock=lambda: next(ticks))
    assert result.status == "deadline" and result.status not in FINAL_STATUSES
    assert "[Page 1]" in result.text and "[Page 5]" not in result.text and result.usable


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Quarterly budget review")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Cost"
    table.cell(1, 0).text = "Servers"
    table.cell(1, 1).text = "1200"
    document.add_paragraph("Closing remarks")
    path = tmp_path / "budget.docx"
    document.save(path)
    result = read_document_text(path)
    assert result.status == "ok" and result.method == "docx"
    assert result.text.split("\n") == ["Quarterly budget review", "Item | Cost", "Servers | 1200", "Closing remarks"]


def test_docx_alternate_content_fallback_is_not_read_twice(tmp_path):
    body = (
        "<w:p><w:r><w:t>Intro line</w:t></w:r></w:p>"
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:p><w:r><w:t>Textbox once</w:t></w:r></w:p></mc:Choice>'
        "<mc:Fallback><w:p><w:r><w:t>Textbox once</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>"
    )
    result = read_document_text(make_docx_package(tmp_path / "boxes.docx", word_document(body)))
    assert result.text.split("\n") == ["Intro line", "Textbox once"]


def test_docx_tabs_and_breaks_are_kept_inside_a_paragraph(tmp_path):
    body = "<w:p><w:r><w:t>Name</w:t><w:tab/><w:t>Value</w:t><w:br/><w:t>Next</w:t></w:r></w:p>"
    result = read_document_text(make_docx_package(tmp_path / "tabs.docx", word_document(body)))
    assert result.text == "Name\tValue\nNext"


def test_text_budget_truncates_and_marks(tmp_path):
    body = "".join(f"<w:p><w:r><w:t>paragraph number {index}</w:t></w:r></w:p>" for index in range(50))
    path = make_docx_package(tmp_path / "many.docx", word_document(body))
    result = read_document_text(path, DocumentReadLimits(max_chars=80))
    assert result.truncated is True and result.status == "ok"
    assert result.text.endswith("[Text truncated at 80 characters]")
    assert len(result.text) < 80 + 60


def test_xlsx_sheets_rows_and_cell_types(tmp_path):
    path = make_xlsx(
        tmp_path / "budget.xlsx",
        [
            ("Q3 Revenue", [["Region", "Total"], ["West", 1200.0], ["East", 980], [None, None], ["Active", True]]),
            ("Notes", [[("inline", "Inline note"), ("error", "#DIV/0!"), 3.5]]),
        ],
    )
    result = read_document_text(path)
    assert result.status == "ok" and result.method == "xlsx" and result.units == 2
    assert result.text.split("\n") == [
        "## Sheet: Q3 Revenue",
        "R1: Region | Total",
        "R2: West | 1200",
        "R3: East | 980",
        "R5: Active | TRUE",
        "## Sheet: Notes",
        "R1: Inline note | 3.5",
    ]


def test_xlsx_phonetic_runs_are_not_added_to_shared_strings(tmp_path):
    path = make_xlsx(tmp_path / "jp.xlsx", [("S", [["plain"]])], phonetic_string=("漢", "かん"))
    members = {name: zipfile.ZipFile(path).read(name) for name in zipfile.ZipFile(path).namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b"<v>0</v></c>", b"<v>0</v></c><c r=\"B1\" t=\"s\"><v>1</v></c>"
    )
    rewritten = write_zip(tmp_path / "jp2.xlsx", members)
    result = read_document_text(rewritten)
    assert "R1: plain | 漢" in result.text and "かん" not in result.text


def test_xlsx_row_cap_truncates(tmp_path):
    rows = [[f"row{index}", index] for index in range(10)]
    result = read_document_text(make_xlsx(tmp_path / "big.xlsx", [("Data", rows)]), DocumentReadLimits(max_rows=3))
    assert result.truncated is True
    assert "row2" in result.text and "row3" not in result.text


def test_xlsx_without_workbook_part_is_an_error(tmp_path):
    result = read_document_text(write_zip(tmp_path / "broken.xlsx", {"docProps/app.xml": "<x/>"}))
    assert result.status == "error"


def test_pptx_follows_presentation_order_and_reads_notes_without_slide_numbers(tmp_path):
    path = make_pptx(
        tmp_path / "deck.pptx",
        [(["First file slide"], None), (["Second file slide", "Bullet two"], ["Speaker note text"])],
        presentation_order=[1, 0],
    )
    result = read_document_text(path)
    assert result.status == "ok" and result.method == "pptx" and result.units == 2
    assert result.text.split("\n") == [
        "## Slide 1",
        "Second file slide",
        "Bullet two",
        "Notes:",
        "Speaker note text",
        "## Slide 2",
        "First file slide",
    ]
    assert "7" not in result.text


def test_pptx_slide_limit_truncates(tmp_path):
    path = make_pptx(tmp_path / "long.pptx", [([f"slide text {index}"], None) for index in range(4)])
    result = read_document_text(path, DocumentReadLimits(max_slides=2))
    assert result.truncated is True and "slide text 1" in result.text and "slide text 2" not in result.text


def test_iwork_reads_embedded_preview_and_says_it_is_partial(tmp_path):
    result = read_document_text(make_iwork(tmp_path / "memo.pages", ["Memo preview words"]))
    assert result.status == "ok" and result.method == "iwork_preview"
    assert "Memo preview words" in result.text and "first page" in result.detail


def test_iwork_without_preview_is_unsupported(tmp_path):
    result = read_document_text(make_iwork(tmp_path / "plain.numbers", None))
    assert result.status == "unsupported" and "preview" in result.detail


@pytest.mark.parametrize("name", ["old.doc", "old.xls", "old.ppt", "notes.txt", "photo.png"])
def test_legacy_and_other_extensions_are_unsupported_without_touching_the_file(tmp_path, name):
    result = read_document_text(tmp_path / name)
    assert result.status == "unsupported" and result.text == ""


def test_signature_mismatch_is_an_error(tmp_path):
    fake_pdf = tmp_path / "fake.pdf"
    fake_pdf.write_text("this is not a pdf", encoding="utf-8")
    fake_docx = tmp_path / "fake.docx"
    fake_docx.write_text("this is not a zip", encoding="utf-8")
    assert read_document_text(fake_pdf).status == "error"
    assert read_document_text(fake_docx).status == "error"


def test_missing_file_and_directory_are_errors_not_exceptions(tmp_path):
    assert read_document_text(tmp_path / "gone.docx").status == "error"
    folder = tmp_path / "folder.pdf"
    folder.mkdir()
    assert read_document_text(folder).status == "error"


def test_symlinked_document_is_refused(tmp_path):
    real = make_pdf(tmp_path / "real.pdf", ["secret outside root"])
    link = tmp_path / "link.pdf"
    os.symlink(real, link)
    result = read_document_text(link)
    assert result.status == "error" and result.text == ""


def test_oversized_file_is_refused_before_parsing(tmp_path):
    path = make_pdf(tmp_path / "plan.pdf", ["words"])
    result = read_document_text(path, DocumentReadLimits(max_bytes=100))
    assert result.status == "too_large"


def test_zip_member_that_expands_past_the_limit_is_refused(tmp_path):
    path = write_zip(tmp_path / "bomb.docx", {"word/document.xml": b"<w:document>" + b" " * 3_000_000 + b"</w:document>"})
    assert path.stat().st_size < 100_000
    result = read_document_text(path, DocumentReadLimits(max_member_bytes=1_000_000))
    assert result.status == "too_large"


def test_xml_entity_declarations_are_rejected(tmp_path):
    xml = (
        '<?xml version="1.0"?><!DOCTYPE w [<!ENTITY boom "expanded">]>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        "<w:p><w:r><w:t>&boom;</w:t></w:r></w:p></w:body></w:document>"
    )
    result = read_document_text(make_docx_package(tmp_path / "entity.docx", xml))
    assert result.status == "error" and "entities" in result.detail and result.text == ""


def test_truncated_xml_keeps_the_text_read_before_the_cut(tmp_path):
    xml = word_document("<w:p><w:r><w:t>kept paragraph</w:t></w:r></w:p><w:p><w:r><w:t>cut")
    xml = xml[: xml.index("cut") + 3]
    result = read_document_text(make_docx_package(tmp_path / "cut.docx", xml))
    assert result.status == "ok" and "kept paragraph" in result.text


def test_is_document_path_is_extension_based_and_case_insensitive():
    assert is_document_path("/tmp/A.PDF") and is_document_path("b.xlsx") and is_document_path("c.pages")
    assert not is_document_path("d.txt") and not is_document_path("e.doc")
